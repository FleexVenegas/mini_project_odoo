import base64
import io

from openpyxl import load_workbook

from odoo import fields, models
from odoo.exceptions import UserError


class SaleOrderImport(models.TransientModel):
    _name = 'sale.order.import'
    _description = 'Import Sale Orders from Excel'

    file = fields.Binary(
        string='Archivo Excel: ',
        required=True,
    )

    filename = fields.Char(
        string='Nombre del Archivo',
    )

    confirm_orders = fields.Boolean(
        string='Confirmar Pedidos de Venta',
        default=False,
        help='Marca esta casilla para confirmar los pedidos de venta.',
    )

    def action_import(self):
        self.ensure_one()

        if not self.file:
            raise UserError('Debes seleccionar un archivo Excel.')

        # ---------------------------------------------------------
        # 1. LEER EXCEL
        # ---------------------------------------------------------

        try:
            file_data = base64.b64decode(self.file)

            workbook = load_workbook(
                filename=io.BytesIO(file_data),
                data_only=True,
            )

        except Exception as error:
            raise UserError(
                f'No se pudo leer el archivo Excel:\n{error}'
            )

        sheet = workbook.active

        rows = list(sheet.iter_rows(values_only=True))

        if not rows:
            raise UserError('El archivo Excel está vacío.')

        # ---------------------------------------------------------
        # 2. LEER ENCABEZADOS
        # ---------------------------------------------------------

        headers = [
            str(header).strip() if header is not None else ''
            for header in rows[0]
        ]

        if len(headers) < 3:
            raise UserError(
                'El Excel debe contener al menos:\n'
                'SKU | Costo | Contacto'
            )

        if headers[0].lower() != 'sku':
            raise UserError(
                'La primera columna debe llamarse "SKU".'
            )

        if headers[1].lower() != 'costo':
            raise UserError(
                'La segunda columna debe llamarse "Costo".'
            )

        partner_names = headers[2:]

        partner_names = [
            name for name in partner_names
            if name
        ]

        if not partner_names:
            raise UserError(
                'Debes indicar al menos un contacto.'
            )

        # Detectar contactos duplicados en encabezados.
        # Sin esta validación, dos columnas con el mismo contacto
        # se fusionarían en una sola cotización sin avisar.
        seen_names = set()
        duplicated_names = set()

        for partner_name in partner_names:
            key = partner_name.lower()

            if key in seen_names:
                duplicated_names.add(partner_name)
            else:
                seen_names.add(key)

        if duplicated_names:
            names_list = ', '.join(sorted(duplicated_names))
            raise UserError(
                f'Los siguientes contactos están duplicados en '
                f'los encabezados:\n{names_list}'
            )

        # ---------------------------------------------------------
        # 3. BUSCAR CONTACTOS
        # ---------------------------------------------------------

        partners = {}

        for partner_name in partner_names:

            partner = self.env['res.partner'].search(
                [('name', '=ilike', partner_name)],
                limit=2,
            )

            if not partner:
                raise UserError(
                    f'No se encontró el contacto:\n'
                    f'"{partner_name}"'
                )

            if len(partner) > 1:
                raise UserError(
                    f'El contacto "{partner_name}" '
                    f'coincide con múltiples contactos en Odoo.'
                )

            partners[partner_name] = partner

        # ---------------------------------------------------------
        # 4. PRECARGAR PRODUCTOS (evita N+1 queries)
        # ---------------------------------------------------------

        product_model = self.env['product.product']

        skus_in_file = set()

        for row in rows[1:]:

            if not any(value is not None for value in row):
                continue

            sku = row[0]

            if sku:
                skus_in_file.add(str(sku).strip())

        all_products = product_model.search(
            [('default_code', 'in', list(skus_in_file))]
        )

        products_by_sku = {}

        for product in all_products:
            products_by_sku.setdefault(
                product.default_code, []
            ).append(product)

        # ---------------------------------------------------------
        # 5. VALIDAR Y PREPARAR LÍNEAS
        # ---------------------------------------------------------

        orders_data = {
            partner.id: {
                'partner': partner,
                'lines': [],
            }
            for partner in partners.values()
        }

        for row_number, row in enumerate(rows[1:], start=2):

            # Ignorar filas completamente vacías
            if not any(value is not None for value in row):
                continue

            sku = row[0]
            base_price = row[1]

            # -----------------------------
            # SKU
            # -----------------------------

            if not sku:
                raise UserError(
                    f'Fila {row_number}: el SKU está vacío.'
                )

            sku = str(sku).strip()

            # -----------------------------
            # PRECIO
            # -----------------------------

            if base_price is None:
                raise UserError(
                    f'Fila {row_number}: '
                    f'el precio base está vacío.'
                )

            try:
                base_price = float(base_price)

            except (TypeError, ValueError):
                raise UserError(
                    f'Fila {row_number}: '
                    f'el precio "{base_price}" no es válido.'
                )

            if base_price < 0:
                raise UserError(
                    f'Fila {row_number}: '
                    f'el precio no puede ser negativo.'
                )

            # -----------------------------
            # PRODUCTO (desde caché precargada)
            # -----------------------------

            matching_products = products_by_sku.get(sku)

            if not matching_products:
                raise UserError(
                    f'Fila {row_number}: '
                    f'no se encontró el SKU "{sku}".'
                )

            if len(matching_products) > 1:
                raise UserError(
                    f'Fila {row_number}: '
                    f'el SKU "{sku}" coincide con múltiples productos.'
                )

            product = matching_products[0]

            # -----------------------------
            # CANTIDADES POR CONTACTO
            # -----------------------------

            for column_index, partner_name in enumerate(
                partner_names,
                start=2,
            ):

                if column_index >= len(row):
                    quantity = None
                else:
                    quantity = row[column_index]

                # Vacío = no cotiza ese producto
                if quantity in (None, ''):
                    continue

                try:
                    quantity = float(quantity)

                except (TypeError, ValueError):
                    raise UserError(
                        f'Fila {row_number}, contacto '
                        f'"{partner_name}": '
                        f'la cantidad "{quantity}" no es válida.'
                    )

                if quantity < 0:
                    raise UserError(
                        f'Fila {row_number}, contacto '
                        f'"{partner_name}": '
                        f'la cantidad no puede ser negativa.'
                    )

                # 0 = no crear línea
                if quantity == 0:
                    continue

                partner = partners[partner_name]

                orders_data[partner.id]['lines'].append({
                    'product': product,
                    'quantity': quantity,
                    'price': base_price,
                })

        # ---------------------------------------------------------
        # 6. CREAR COTIZACIONES (batch create)
        # ---------------------------------------------------------

        SaleOrder = self.env['sale.order']

        orders_to_create = []

        for data in orders_data.values():

            # No crear cotización si no tiene productos
            if not data['lines']:
                continue

            order_lines = []

            for line in data['lines']:

                order_lines.append(
                    (
                        0,
                        0,
                        {
                            'product_id': line['product'].id,
                            'product_uom_qty': line['quantity'],
                            'price_unit': line['price'],
                        },
                    )
                )

            orders_to_create.append({
                'partner_id': data['partner'].id,
                'order_line': order_lines,
            })

        if not orders_to_create:
            raise UserError(
                'No se creó ninguna cotización. '
                'Todas las cantidades están vacías o en cero.'
            )

        created_orders = SaleOrder.create(orders_to_create)

        # ---------------------------------------------------------
        # 7. RESULTADO
        # ---------------------------------------------------------

        order_names = ', '.join(created_orders.mapped('name'))


        # ---------------------------------------------------------
        # ---------------------------------------------------------
        # 8. CONFIRMAR COTIZACIONES SI ESTÁ ACTIVADA LA BANDERA
        # ---------------------------------------------------------
        if self.confirm_orders:
            created_orders.action_confirm()


        return {
            'type': 'ir.actions.client',
            'tag': 'delayed_view_reload',
            'params': {
                'title': 'Importación completada',
                'message': (
                    f'Se crearon {len(created_orders)} cotización(es): '
                    f'{order_names}'
                ),
                'type': 'success',
                'sticky': False,
                'delay': 1500,
            },
        }